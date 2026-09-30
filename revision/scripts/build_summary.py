"""Write the requested short, high-level author report."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
ROOT=Path(__file__).resolve().parents[1]
SECTIONS=[
('Changes made',[
'Updated phonic richness results and figures. The positive contrasts across all projects remained when Indonesian recordings were subsampled to match recording effort elsewhere.',
'Changed grazing and shrimp models to raw detections with recorded-effort offsets; corrected the cuescape Methods to match its existing exposure adjustment. Fixed a Results error attributing the large Mexican shrimp contrast to healthy rather than mid-stage restored habitat.',
'Added Indonesia-exclusion and alternative-model checks. Removing Indonesia substantially weakens the evidence: no remaining habitat contrast against degraded reefs excludes zero at 95% confidence. The draft now qualifies generalisation.',
'Added clearer limits on causal interpretation, acoustic proxies, baseline selection and single-site restoration comparisons. Shortened Text S1 to explain why each check was done, how it was done and what it found. Figures retain their original style and both 75% and 95% intervals, with updated model values.'
]),
('Tested or considered, but not adopted as firm conclusions',[
'No claim that country effects eliminate Indonesia’s influence. No global fish-richness/biomass gradient was invented from missing survey data.',
'Age and cover analyses remain exploratory. Six Indonesian sites supply the within-country age information; site-level age tests are inconclusive. Coral cover changes some contrasts but may itself be part of restoration’s effect. Reef-star intensity cannot be separated reliably from broader project-level differences.',
'Thesis material was used selectively; uncertain factual changes and duplicated speculative discussion were not imported.'
]),
('Your remaining decisions',[
'Review the proposed applied Abstract/Discussion wording, statistical adequacy, Mexico site identity, and support for interpreting snapping shrimp as bioturbation. Finalise anonymous data availability and submission details. All are linked to numbered comments, current lines and suggested wording in the response document’s bottom TODO section.'
])]

def main() -> None:
    """Produce matching one-page-oriented DOCX and Markdown reports."""
    d=Document();d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10)
    d.styles['Normal'].paragraph_format.space_after=Pt(4)
    d.add_heading('Journal revision — first-pass summary',0)
    md=['# Journal revision — first-pass summary','']
    for title,points in SECTIONS:
        d.add_heading(title,2);md.extend(['## '+title,''])
        for text in points:d.add_paragraph(text,style='List Bullet');md.append('- '+text)
        md.append('')
    footer='Original documents, code and data were preserved. Clean and tracked files, concise responses and reproducible analyses are supplied; author decisions remain before submission.'
    d.add_paragraph(footer);md.append(footer)
    d.core_properties.author='';d.core_properties.last_modified_by=''
    d.save(ROOT/'outputs/Summary_report.docx')
    (ROOT/'outputs/Summary_report.md').write_text('\n'.join(md)+'\n')

if __name__=='__main__':main()
