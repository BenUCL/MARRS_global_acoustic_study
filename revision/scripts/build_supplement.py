"""Append a concise, tracked sensitivity appendix to the original supplement."""
from copy import deepcopy
import sys
from pathlib import Path
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt
from lxml import etree
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from docx_revision import RevisionDocument,tag,NS
METRICS={'settlement_cuescape':'Recruitment cuescape','phonic_richness':'Phonic richness','graze_count':'Grazing','snaps_count':'Shrimp activity'}
SCENARIOS={'settlement_cuescape':'baseline','phonic_richness':'corrected_baseline','graze_count':'actual_effort','snaps_count':'actual_effort'}

def main() -> None:
    """Preserve all existing tables, images and references; append analysis results."""
    models=pd.read_csv(ROOT/'outputs/all_model_coefficients.csv')
    d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
    d.add_heading('Text S3. Sensitivity analyses',level=1)
    paragraphs=[
        'Finalised models: The following checks assess the finalised models (Table S3, “All projects”; sample sizes in Table S6). To account for unequal recording time, grazing and shrimp counts were modelled as negative binomial counts with log(recorded hours) as an offset, estimating detections per hour of recording time available. Recruitment cuescape models used log(recorded nighttime hours) as the offset. These measure detection rates, not animal abundance. Grazing models excluded Kenya, where the grazing sonotype was not detected.',
        'Model assumptions: Richness varied less than the negative binomial model assumes. The date-within-project variance was effectively zero in the finalised phonic-richness model (2.62 × 10⁻¹³) and zero in the richness model with fixed project effects. In the finalised shrimp model, date-within-project variance was small (0.0031), but did not collapse to zero. We therefore retained these date effects because our model structure was ecologically motivated; recordings from the same project and date may share environmental conditions. Date effects accounted for variation in several other models, so retaining them provided a consistent approach rather than selecting them separately for each indicator. To test alternative ways of accounting for project and day-to-day differences, we fitted project as a fixed rather than random effect and added a date effect to the pooled cuescape model. For both checks, every habitat comparison against degraded reefs kept the same direction. Significant differences remained significant and non-significant differences remained non-significant at p < 0.05 (Table S7).',
        'Indonesia’s influence: To test whether Indonesia’s larger sample (21 of 45 sites; Table S6) strongly influenced pooled findings, we refitted all four indicators without it. Every habitat comparison against degraded reefs kept the same direction, but previously significant differences became non-significant; the other comparisons remained non-significant. All 95% confidence intervals included no difference (Table S3). Including Indonesia adds site replication, which can improve precision and increase the chance of detecting real habitat differences as statistically significant. This check cannot separate that contribution from project-specific differences.',
        'Recording effort and richness: To test whether longer recording time increased observed richness, we reduced Indonesia’s schedule from one minute every two minutes to one every four, matching other projects. We tested both alternate-recording subsets separately, retaining whole one-minute recordings. Days required at least 324 of 360 scheduled recordings; both comparisons retained the same 1,060 site-days. After recalculating richness and refitting the pooled model, all three habitat comparisons against degraded reefs kept the same positive direction and remained statistically significant in both subsets. Thus, downsampling did not change either the direction or significance of these comparisons (all p < 0.03; Table S4).',
        'Coral cover: To test whether habitat comparisons changed after accounting for coral cover, models with and without cover used the same observations from sites with matched cover measurements. Values reported as <5% were tested as 0%, 2.5% and 5%. Adding cover weakened some cuescape comparisons; the mid-stage grazing comparison was positive (p = 0.018) under the 2.5% assumption (Table S8). These checks are exploratory: cover may itself respond to restoration, so adjustment changes the question being tested.',
    ]
    for p in paragraphs:d.add_paragraph(p)
    def table(title,headers,rows):
        d.add_heading(title,level=2);t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        for row in rows:
            for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
        for row in t.rows:
            properties=row._tr.get_or_add_trPr()
            etree.SubElement(properties,tag('cantSplit'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1
                    p.paragraph_format.space_before=Pt(0)
                    p.paragraph_format.space_after=Pt(4)
                    for r in p.runs:r.font.size=Pt(9)
        etree.SubElement(t.rows[0]._tr.get_or_add_trPr(),tag('tblHeader'))
    def result(metric,scenario,term):
        r=models[(models.metric==metric)&(models.scenario==scenario)&(models.term==term)].iloc[0]
        return f'{r.estimate:.2f} [{r.lower95:.2f}, {r.upper95:.2f}]; p={r.p:.3g}'
    rows=[]
    for metric,label in METRICS.items():
        ex={'phonic_richness':'exclude_indonesia','graze_count':'actual_effort_exclude_indonesia','snaps_count':'actual_effort_exclude_indonesia'}.get(metric,'exclude_indonesia')
        for term,habitat in [('treatmenthealthy','Healthy'),('treatmentnewly_restored','Early-stage'),('treatmentrestored','Mid-stage')]:
            rows.append([label,habitat,result(metric,SCENARIOS[metric],term),result(metric,ex,term)])
    table('Table S3. Pooled log response ratios [95% CI] against degraded reefs',['Indicator','Habitat','All projects','Without Indonesia'],rows)
    d.add_paragraph('Tables S3–S4 and S7–S8 report log response ratios (zero = no difference; positive = higher than degraded), 95% confidence intervals and two-sided Wald p-values without adjustment for multiple comparisons.')
    rows=[]
    for term,habitat in [('treatmenthealthy','Healthy'),('treatmentnewly_restored','Early-stage'),('treatmentrestored','Mid-stage')]:
        rows.append([habitat]+[result('phonic_richness',{'reconstructed':'corrected_baseline','phase0':'duty_phase0','phase2':'duty_phase2'}[phase],term)
                              for phase in ['reconstructed','phase0','phase2']])
    table('Table S4. Phonic richness under full and matched recording effort',
          ['Habitat','Full effort','Four-minute subset 1','Four-minute subset 2'],rows)
    ages=pd.read_csv(ROOT/'outputs/indonesia_age_site_permutation.csv')
    table('Table S5. Site-level age associations within Indonesia',['Indicator','Independent sites','Spearman ρ','Exact two-sided p'],[[METRICS[r.metric],r.sites,f'{r.rho:.3f}',f'{r.exact_two_sided_p:.3f}'] for r in ages.itertuples()])
    d.add_paragraph('Spearman ρ measures whether indicators increase (positive) or decrease (negative) with age. Two-sided p-values use all possible age assignments among the six sites.')
    counts=[]
    for metric,label in METRICS.items():
        data=pd.read_csv(ROOT/'data/results/functions'/f'{metric}.csv')
        if metric=='graze_count':data=data[data.country!='kenya']
        for (country,treatment),g in data.groupby(['country','treatment']):counts.append(dict(indicator=label,country=country,habitat=treatment,sites=g.site.nunique(),days=len(g)))
    counts=pd.DataFrame(counts);counts.to_csv(ROOT/'outputs/sample_sizes.csv',index=False)
    pivot=counts.pivot(index=['country','habitat'],columns='indicator',values='days').fillna(0).astype(int)
    sites=counts.groupby(['country','habitat']).sites.max()
    table('Table S6. Spatial replication and recording-day observations',['Project','Habitat','Sites','Cuescape days','Richness days','Grazing days','Shrimp days'],[[country.capitalize(),{'newly_restored':'Early-stage','restored':'Mid-stage','healthy':'Healthy','degraded':'Degraded'}[habitat],sites.loc[(country,habitat)],r['Recruitment cuescape'],r['Phonic richness'],r['Grazing'],r['Shrimp activity']] for (country,habitat),r in pivot.iterrows()])
    # Keep existing table numbers stable; report comparisons from stored fits.
    habitats=[('treatmenthealthy','Healthy'),('treatmentnewly_restored','Early-stage'),('treatmentrestored','Mid-stage')]
    rows=[]
    for metric,label in METRICS.items():
        for term,habitat in habitats:
            rows.append([label,habitat,result(metric,SCENARIOS[metric],term),
                         result(metric,'country_fixed',term),
                         result(metric,'include_date',term) if metric=='settlement_cuescape' else '—'])
    table('Table S7. Habitat comparisons under alternative model assumptions',
          ['Indicator','Habitat','Finalised model','Fixed project effects','Added date effect'],rows)
    d.add_paragraph('Each alternative changes the feature named in its column. All indicators use negative binomial models. The date check adds a random date intercept to the pooled cuescape model. Dashes indicate checks that do not apply.')
    rows=[]
    for metric,label in METRICS.items():
        for term,habitat in habitats:
            rows.append([label,habitat]+[result(metric,scenario,term) for scenario in
                         ['cover_complete_baseline','cover_bound_0','cover_bound_2.5','cover_bound_5']])
    table('Table S8. Habitat comparisons with and without adjustment for coral cover',
          ['Indicator','Habitat','Without cover','With cover: <5% as 0%','With cover: <5% as 2.5%','With cover: <5% as 5%'],rows)
    d.add_paragraph('Within each indicator, all four models use the same observations with matched coral-cover measurements. The three adjusted models differ only in the value assigned to cover reported as <5%. Compare columns within this table rather than attributing differences from Table S3 solely to coral cover.')
    age = pd.read_csv(ROOT / 'data/results/functions/sites_age.csv')
    age_rows = age[age.site.str.startswith('R')].copy()
    site_values = age_rows[['country', 'site', 'age']].copy()
    for metric, column in [
        ('phonic_richness', 'Phonic richness\n(mean sound types/day)'),
        ('settlement_cuescape', 'Recruitment cuescape\n(mean detections/nighttime hour)'),
        ('graze_count', 'Grazing\n(mean detections/hour)'),
        ('snaps_count', 'Shrimp activity\n(mean detections/hour)'),
    ]:
        if metric == 'phonic_richness':
            values = pd.read_csv(ROOT / 'data/richness_reconstructed.csv')
        elif metric in ['graze_count', 'snaps_count']:
            values = pd.read_csv(ROOT / 'data' / f'{metric}_effort.csv')
            values['count'] /= values.recorded_hours
        else:
            values = pd.read_csv(ROOT / 'data/results/functions/settlement_cuescape.csv')
            values['count'] /= values.max_poss_count / 720
        values = values.groupby(['country', 'site'], as_index=False)['count'].mean()
        site_values = site_values.merge(values, on=['country', 'site'], how='left').rename(columns={'count': column})
    site_values = site_values.sort_values('age')
    table('Table S9. Mean acoustic indicator values at mid-stage restoration sites',
          ['Project', 'Site', 'Restoration age\n(months)', 'Phonic richness\n(mean sound types/day)',
           'Recruitment cuescape\n(mean detections/nighttime hour)', 'Grazing\n(mean detections/hour)',
           'Shrimp activity\n(mean detections/hour)'],
          [[row.country.capitalize(), row.site, row.age, f'{row[3]:.2f}', f'{row[4]:.2f}',
            f'{row[5]:.2f}', f'{row[6]:.2f}'] for row in site_values.itertuples(index=False)])
    d.add_paragraph('Values are site means across recording days, rows are ranked by age of restoration. The eight sites do not show a clear monotonic pattern with restoration age; their limited spatial replication prevents a reliable test of age effects.')
    revision=RevisionDocument(ROOT/'docs/Supplementary 1.docx')
    # Minimal grammatical correction, retaining hyperlinks and citations.
    revision.replace(4,'with the Python’s Contextily package (v1.6.2). can be found below.','with Python’s Contextily package (v1.6.2) can be found below.','6','Grammar: remove stray full stop and article')
    revision.replace(4,'mapping software’s','mapping software','6','Grammar: remove incorrect possessive')
    revision.add_comment(4,'Grammar: removed a stray full stop, unnecessary article and incorrect possessive; links preserved.')
    body=revision.root.find(tag('body'));section=body.find(tag('sectPr'))
    for element in d._element.body:
        if element.tag==tag('sectPr'):continue
        block=deepcopy(element)
        for run in list(block.xpath('.//w:r')):
            parent=run.getparent();idx=parent.index(run);parent.remove(run);parent.insert(idx,revision.wrapper('ins',run))
        if block.tag==tag('p'):
            pr=block.find(tag('pPr'))
            if pr is None:pr=etree.Element(tag('pPr'));block.insert(0,pr)
            rp=etree.SubElement(pr,tag('rPr'))
            marker=revision.wrapper('ins',etree.Element(tag('r')));marker.remove(marker[0]);rp.append(marker)
        if block.tag==tag('tbl'):
            for row in block.findall(tag('tr')):
                pr=row.find(tag('trPr'))
                if pr is None:pr=etree.SubElement(row,tag('trPr'))
                marker=revision.wrapper('ins',etree.Element(tag('r')));marker.remove(marker[0]);pr.append(marker)
        body.insert(body.index(section) if section is not None else len(body),block)
    revision.write(ROOT/'outputs/Supplementary_1_tracked.docx')
    revision.write(ROOT/'outputs/Supplementary_1_clean.docx',clean=True)
    print('Supplementary appendix written with 7 results tables.')

if __name__=='__main__':main()
