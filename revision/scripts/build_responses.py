"""Build reference and populated response documents with traceable author actions."""
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
ROOT=Path(__file__).resolve().parents[1]
BLUE=RGBColor.from_string('245A81')
REGISTER=json.loads((ROOT/'outputs/comment_register.json').read_text())
LINES=json.loads((ROOT/'outputs/line_map.json').read_text())
RESPONSES={
1:('Implemented',[5,62],'Thank you. We have replaced “significantly enhanced” with wording describing differences from degraded baselines and have clarified that the cross-sectional comparisons do not establish causality.'),
2:('Implemented; interpretation requires author review',[10,67],'We now define bioturbation as the mixing and reworking of sediment by organisms and explain its relevance to nutrient cycling. We also make clear that shrimp snaps are not a calibrated measure of sediment turnover.'),
3:('Implemented',[12],'We have briefly explained that rubble stabilisation secures loose coral rubble to provide a stable substrate for coral fragments, with a pointer to section 3.1.'),
4:('Clarified',[20],'The recorders were HydroMoths, as stated at the start of this paragraph. AudioMoth is the name of the configuration application. We have clarified this distinction while retaining the correct software name.'),
5:('Implemented',[26],'We now define embeddings as numerical representations that summarise acoustic features for comparison and classification when the term first appears.'),
6:('Implemented',[35],'We have changed “used to inference” to “used to run inference”. Minor grammatical corrections are also marked with tracked changes and explanatory comments.'),
7:('Implemented', [39, 40, 46], 'We agree that detection counts do not estimate numbers of individual fish. The cuescape, grazing and shrimp models now use raw detection counts with the log of recorded hours as an exposure offset; for the cuescape, these are nighttime hours only. We explicitly distinguish acoustic detection rates from animal abundance. These did not change the conclusions; further details are now shared in Supp. Text S3 under the “Finalised models” paragraph, with a deeper exploration for phonic richness in the “Recording effort and richness” paragraph.'),
8:('Implemented', [40], 'This is addressed in response to comment 7. Shrimp counts are now modelled with log(recorded hours) as an exposure offset. The pooled habitat comparisons against degraded reefs retained the same directions and significance decisions: significant differences remained significant and non-significant differences remained non-significant. See Supp. Text S3, “Finalised models”, with the finalised estimates in Table S3.'),
9:('Implemented', [46, 50, 52, 59], 'The thick 75% and thin 95% intervals are retained as complementary descriptions of uncertainty. Statistical significance is assessed using two-sided Wald tests and the 95% interval against zero percentage change, not the 75% interval or overlap between separate habitat intervals. We report non-significant contrasts as such, identify the tests as unadjusted, and no longer treat a non-significant healthy–restored comparison as evidence of equivalence. See Supp. Text S3, “Finalised models”, and the statistical reporting note accompanying Tables S3–S4 for the estimates and test conventions.'),
10:('Analyses completed; broader conclusion wording for author review', [50, 62], 'We agree that project random effects do not rule out disproportionate influence. Indonesia contributed 21 of the 45 sites. We refitted all four indicators without Indonesia and describe this check in Supp. Text S3, “Indonesia’s influence”, with estimates, 95% intervals and p-values in Table S3. Every habitat comparison against degraded reefs kept the same direction, but previously significant differences became non-significant; the other comparisons remained non-significant. All 95% confidence intervals included zero. For example, the mid-stage cuescape coefficient changed from 1.61 (p < 0.001) to 1.06 (p = 0.212). Including Indonesia adds site replication, which can improve precision and increase the chance of detecting real habitat differences as statistically significant. This check cannot separate that contribution from project-specific differences. We therefore retain the qualification on generalisation across regions.'),
11:('Implemented and explored', [17, 67], 'These boundaries describe the sampled age groups, not established recovery stages. We have clarified this and corrected early-stage to ≤3 months to match Table S1. We agree that age relationships would be informative. Table S9 presents the restoration age and mean acoustic indicator values for all eight mid-stage sites. It shows no clear monotonic pattern, but eight sites do not provide sufficient spatial replication to determine whether age effects are absent or to estimate recovery rates. We therefore identify longer-term repeated sampling as the appropriate way to examine recovery trajectories.'),
12:('Implemented', [39, 44, 50], 'Phonic richness was not multiplied by duty cycle, but we agree that additional recording time may increase observed richness. To test this, we reduced the Indonesian recordings from one minute every two minutes to one every four, matching other projects. We tested both alternate-recording subsets separately, recalculated daily richness and refitted the pooled model using the same site-days. All three pooled habitat comparisons against degraded reefs kept the same positive direction and remained statistically significant in both subsamples. Downsampling therefore changed neither the direction nor the significance of these comparisons (all p < 0.03; Table S4). See Supp. Text S3, “Recording effort and richness”, and Table S4. The finalised models retain full recording effort; the reduced-effort datasets are sensitivity checks.'),
13:('Explored; limitations documented', [46], 'We explored coral cover by comparing models with and without it on the same observations, as described in Supp. Text S3, “Coral cover”, and Table S8. Across all three assumptions for cover reported as <5%, the mid-stage grazing comparison against degraded reefs changed from non-significant to significantly positive. The positive healthy and mid-stage cuescape comparisons, mid-stage richness comparison and healthy shrimp comparison became non-significant. The healthy grazing estimate changed from negative to positive but remained non-significant; all other directions and significance decisions were unchanged. Cover may itself respond to restoration, so these are exploratory comparisons after accounting for cover, rather than replacements for the finalised habitat models. Fish biomass and richness were available only in the Maldives. Reef-star numbers varied mainly among projects, with insufficient variation among sites within the same project to separate restoration intensity reliably from broader project-level differences. One Mexico site-label discrepancy remains for author confirmation.'),
14:('Implemented',[17,65],'We have clarified which survey measures supported site selection. In the Maldives, fish biomass and richness informed the baseline categories; the acoustic contrasts are therefore not an independent validation of those attributes. This limitation is now stated in the Discussion.'),
15:('Limitation implemented; current stressor details require author input',[65],'We now carry the local-reference caveat into the Discussion and distinguish the reported causes of degradation in Table S1 from ongoing stressor exposure, which was not quantified at the reference sites. We cannot infer current exposure from a historical cause of degradation. A concise project-specific description can be added if verified observations are available.'),
16:('Implemented; error-source explanation requires author input',[36,40,67],'We now state that the Scrape attribution has not been validated by direct observation in every region. We explicitly report that 26 of 100 checked Mexican detections were false positives and acknowledge that this uncertainty may obscure habitat differences. This is a false-positive proportion among retained detections, not the conventional false-positive rate among negatives. Without reviewing the misclassified recordings, we cannot identify the alternative sound sources or explain regional differences reliably.'),
17:('Implemented', [56, 67], 'We now identify Kenya, the Maldives and Australia as single-restored-site comparisons and describe these findings as preliminary. Repeated days do not provide independent spatial replication. The numbers of sites and recording days are given in Table S6, referenced from Supp. Text S3, “Finalised models”.'),
18:('Implemented; optional artwork tidy-up',[25,26,34,36],'We checked the supplied Figure 2 directly. The raw-data box is 1a, not 1c; we have corrected the text and aligned the evaluation references with labels 9–11. The caption now explains the parallel a/b/c branches and the distinction between 13a and 13b. A purely sequential redesign is optional and would require coordinated changes to the artwork and text.'),
19:('Analyses completed; author checks identified', [46, 50, 62, 65, 67], 'Thank you. Indonesia’s influence is addressed in response to comment 10, covariate availability and interpretation in comment 13, and limited spatial replication in comment 17. We also fitted country as a fixed effect and added a date effect to the pooled cuescape model. These checks are described in Supp. Text S3, “Model assumptions”, and Table S7. In both checks, every habitat comparison against degraded reefs kept the same direction: significant differences remained significant and non-significant differences remained non-significant at p < 0.05 (Table S7). The date-within-project variance was effectively zero in the finalised richness model and zero in the fixed-project richness model; it was small (0.0031) but not singular in the finalised shrimp model. We retained these date effects because shared environmental conditions within project and date informed the model structure, date effects accounted for variation in several other models, and a consistent structure avoided selecting terms separately for each indicator. Country effects account for project-level baseline differences, but they neither identify the ecological causes of those differences nor replace site-level measurements. The available surveys do not provide a comparable richness gradient across all five regions, so we have not constructed one from assumed biogeographic rankings.'),
20:('Implemented, subject to final model review', [44, 46, 50, 67], 'We have added the modelling checks to Supp. Text S3 and Tables S3–S9. “Recording effort and richness” (Table S4) and “Model assumptions” (Table S7) show that the recording-effort and alternative project/date checks changed neither the direction nor the significance of any habitat comparison against degraded reefs. “Indonesia’s influence” (Table S3) shows unchanged directions but loss of the previously significant differences. “Coral cover” (Table S8) identifies the comparisons whose direction or significance changed after adjustment. Table S9 presents the eight mid-stage sites’ ages and mean indicator values; it shows no clear monotonic pattern but limited spatial replication. The finalised methods and sample sizes are signposted in “Finalised models”.'),
21:('Proposed wording only, awaiting author review',[],'Suggested Abstract and Discussion wording is provided in the author TODO section below. These interpretation-led edits have not been applied to the manuscript, pending the lead author’s review.'),
22:('Preparation completed in part; submission actions remain',[],'Clean and tracked manuscript and supplement copies are supplied. The fixed line-reference PDF underpins the updated line annotations. Before submission, resolve the author TODOs, remove this internal TODO section from the response, finalise the numbered Abstract and anonymous data-availability statement, check the graphical abstract, and confirm the revision deadline in the submission system.')
}
TASKS=[
('T1','1, 2, 3, 4, 5, 6','Clarify terminology, remove causal wording at the flagged passage and correct grammar.','Implemented'),
('T2','7, 8, 9','Use effort-standardised detections, clarify inference, and update model values in Figures 4–5 while retaining their original style and 75%/95% intervals.','Implemented'),
('T3','10, 19, 20','Test Indonesia exclusion and alternative country/date specifications; qualify pooled findings.','Completed'),
('T4','11','Explain age categories and test continuous age without treating recording days as independent sites.','Completed; exploratory'),
('T5','12','Test whether the higher recording effort in Indonesia affects richness comparisons.','Completed'),
('T6','13, 19','Assess cover, fish metrics and reef-star intensity using matched subsets and report limitations.','Completed; site identity remains'),
('T7','14, 15, 16, 17','Clarify reference-site selection, ongoing-stressor uncertainty, Scrape uncertainty and single-site inference.','Implemented; factual follow-up below'),
('T8','18','Align Figure 2 references and explain branch numbering.','Implemented; optional artwork tidy-up'),
('T9','21, 22','Finalise applied implications, Abstract structure, data statement and submission package.','Author review required')
]
TODO=[
('A1','21, 22c; also 1 and 10',[5,66,70],
 'Make applied value explicit and finish the numbered Abstract; these are suggestions, not applied changes.',
 'Possible first point: “Coral reef restoration is expanding, but managers rarely have evidence of whether ecological processes recover alongside coral cover. Monitoring these outcomes is needed to judge when restoration delivers broader benefits.”\nPossible final point: “Synthesis and applications. Combining passive acoustic monitoring with targeted classifiers can help managers compare restored reefs with local healthy and degraded references. Acoustic indicators should complement ecological surveys, with repeated monitoring and replicated sites used to distinguish local responses from wider recovery.”\nPossible Discussion addition: “Managers could use this approach alongside coral and fish surveys to track which aspects of the soundscape approach local reference conditions, and which remain distinct. These comparisons should inform site-specific monitoring rather than a universal acoustic threshold for restoration success.”\nSplit the remaining Abstract into numbered methods, results and synthesis points after agreeing the claims. Qualify “full recovery” in the Conclusion, for example: “We found acoustic evidence consistent with recovery of some ecological processes following restoration, with substantial variation among projects.”'),
('A2','10, 12, 13, 19, 20',[44,46,50],
 'Review the revised modelling choices with the statistical co-author before treating the draft as submission-ready.',
 'The draft uses negative binomial models for all four indicators; detection counts use actual recorded exposure. Richness is bounded and strongly underdispersed, and several random-effect estimates are on the boundary. Review adequacy, uncertainty and the unadjusted habitat comparisons, including the single-restored-site country models; do not use convergence alone as evidence of fit. If retained, the current Methods wording is provided in the manuscript and Text S3. The pooled cuescape’s original omission of date is now explicit; the full-date sensitivity is available. Any model change must update Results, Figures 4–5 and responses together.'),
('A3','13, 19; also 11',[17],
 'Resolve the Mexico site identity before adding its covariates or altering Table S1.',
 'Table S1 labels the mid-stage site R2 (350 reef stars, 35 months), whereas the analysis CSVs label the mid-stage site R1 and the age CSV gives 35 months. Confirm whether these are the same site. If confirmed, change Table S1 “R2” to “R1” and rerun the covariate analyses with the matched metadata. Also check the manuscript’s 157-star minimum for mid-stage plots, which is not represented in the supplied Table S1. Do not assume these corrections from similar ages alone.'),
('A4','15, 19',[65],
 'Reviewer asks for current stressors at healthy references; Table S1 lists historical degradation causes.',
 'Confirm any documented ongoing fishing, sedimentation, pollution, bleaching or other exposure for each project at the time of recording. Possible sentence, only after filling from records: “During acoustic sampling, the principal ongoing pressures were [verified pressures by project]; healthy references were the least disturbed available sites rather than stressor-free controls.” If those records do not exist, retain the implemented statement that ongoing exposure was not quantified.'),
('A5','16',[40],
 'Reviewer asks why the Mexican Scrape classifier produced 26% false positives and whether other sources varied by region.',
 'Review the 26 false-positive examples or the original annotation notes. Possible addition, only if verified: “False positives were mainly [identified sound/source], which overlapped with Scrape in [documented acoustic feature].” The first pass deliberately does not invent sources or claim classifier difficulty. Audio may need read-only inspection for this item; the numerical revisions used only CSVs.'),
('A6','2, 16; interpretation relevant to 21',[10,64,67],
 'Check whether the cited bioturbation literature supports the snapping-shrimp claim.',
 'Suchanek et al. (1986) concerns callianassid ghost shrimp; Laverock et al. (2010) studied Upogebia deltaura and Callianassa subterranea, rather than the detected Alpheidae. These citations do not validate a relationship between snap rates and sediment turnover. Suggested wording: “Snapping shrimp activity provides an additional measure of benthic acoustic activity; its relationship with sediment bioturbation remains to be established.” Decide whether to adopt this throughout the Abstract, Introduction and Discussion or provide a directly relevant validation source. The manuscript already notes that turnover is not calibrated. Source: https://doi.org/10.1038/ismej.2010.86.'),
('A7','18; also 9',[25,52,59],
 'Optional PowerPoint cleanup and final figure checks.',
 'Figure 2’s existing labels now match the manuscript; a redraw could rename the unsupervised branch and make 13a/13b more intuitive, but all text references must then change together. Figures 4–5 retain their original colours, ecosystem-process axes, percentage scale, N annotations, habitat legend and 75%/95% intervals, with the revised model values. Vector PDFs and high-resolution PNGs are supplied in outputs/figures. Check that Figure 1 shows 21 Indonesian sites and that the graphical abstract reflects the revised, more cautious conclusions.'),
('A8','22a, 22b, 22d, 22e, 22f, 22h',[70],
 'Complete submission administration and anonymity checks.',
 'Add an anonymous data-availability statement before References. Suggested wording: “The data and code supporting this study are available to reviewers at [anonymous repository link]; the final archive will include the revised analysis.” Replace the placeholder with a verified anonymous access mechanism. Check whether the existing supplementary public repository and map links reveal authorship. Confirm the portal deadline or extension: the letter is dated 18 June 2026 and requested return within 90 days. Confirm the graphical abstract meets the journal’s requirements and co-authors can grant the chosen publication licence. An additional language abstract is optional. Recheck final line numbers in Word, word count and any reference fields after your edits, and remove this internal author TODO section before submitting the response.'),
('A9','5, 18; thesis comparison',[16,26,36],
 'Confirm factual differences between thesis chapter 5 and the manuscript; audit bibliography.',
 'The thesis changes frame diameter (0.91 to 0.93 m), coral-fragment size (approximately 10 to 15 cm), maintenance visits and some pipeline labels. These were not imported without factual confirmation. Its direct-validation paragraph informed the added limitation; duplicated speculative explanations were omitted. Navine et al. (2024) is cited in the manuscript but missing from its bibliography; verify the intended source before adding it. Two Tricas and Boyle (2021) bibliography entries appear to describe the same DOI; reconcile using the original reference manager to preserve field links. Do not blindly update citation fields in the edited copies.')
]

def span(paragraphs: list[int]) -> str:
    bits=[]
    for p in paragraphs:
        v=LINES['paragraphs'].get(str(p))
        if v:bits.append(f'{v[0]}–{v[1]}')
    return ', '.join(bits) if bits else 'No manuscript edit yet'

# These point to the actual edited lines in the clean manuscript's fixed PDF.
# A paragraph-wide map is too broad for point-by-point review.
MANUSCRIPT_CHANGES={
1:'Causal wording in the Abstract, line 21; interpretation in the Discussion, lines 527–530.',
2:'Bioturbation definition, lines 83–84; acoustic-proxy limitation, lines 637–640.',
3:'Rubble stabilisation explanation, lines 108–110.',
4:'Recorder and configuration-application distinction, lines 177–178.',
5:'First definition of embeddings, lines 225–226.',
6:'Inference wording, line 297.',
7:'Cuescape detection rate, lines 340–343; grazing and shrimp exposure, lines 359–368; model offset, lines 406–408.',
8:'Shrimp exposure, clean manuscript lines 367–368. The effort model is described in Supp. Text S3, “Finalised models”.',
9:'Test convention, lines 411–413; interpretation of non-significance, line 431.',
10:'Indonesia-exclusion result, lines 446–448; qualified interpretation, lines 527–530.',
12:'Recording-effort check, lines 349–352; result, lines 448–450. See Supp. Text S3, “Recording effort and richness”, Table S4.',
14:'Site-selection survey measures, lines 147–149; limitation, lines 586–590.',
17:'Single-site interpretation, lines 641–644.',
18:'Figure 2 caption, lines 207–211; matching workflow references, lines 222–224 and 295–300.',
20:'Sensitivity methods, lines 395–399 and 413–417; Indonesia result, lines 446–448; recording-effort result, lines 448–450; age limitation, lines 629–632.',
21:'Suggested Abstract and Discussion wording is in the author TODO section; no manuscript edit has been made for this point.',
}

PART_RESPONSES={
11:[
('The manuscript now says that the age groups describe the sampled sites, not established recovery stages, and includes the three-month boundary (clean manuscript lines 149–151, 160–162).', 'Age-group explanation and boundary, clean manuscript lines 149–151 and 160–162.'),
('We agree that age relationships would be informative. Table S9 gives age and site-mean values for all eight mid-stage sites. Sites aged 1–3 months were excluded because the absence of sampled sites between 3 and 32 months means that combining both restoration stages would impose an assumed continuous relationship across a large unsampled age gap. We find there is no clear monotonic pattern. However, given that we are only able to report on eight mid-stage sites, the sample size is very small, making it difficult to establish whether age effects are truly absent or the analysis is statistically underpowered. We now say this alongside the need for long-term monitoring (clean manuscript lines 629–632).', 'Age finding and spatial-replication qualification, clean manuscript lines 629–632; raw values in Supplementary Table S9.'),
],
13:[
('We agree that this would be an informative analysis. However, closer inspection of Supplementary Table S1 shows that reef-star numbers vary mainly among projects, with insufficient variation among sites within the same project. Their effect therefore cannot be reliably separated from broader differences among projects. We now identify restoration intensity as an important question and explain why it could not be tested using the available data (clean manuscript lines 552–556).', 'Restoration-intensity limitation, clean manuscript lines 552–556.'),
('We compared habitat models with and without coral cover on the same sites, testing three plausible values for cover reported as below 5%. Some directions and significance decisions changed, so the cover results remain exploratory (Supp. Text S3, “Coral cover”, Table S8). Fish biomass and richness were available only in the Maldives and could not be assessed as comparable global covariates.', 'Site-covariate analyses, clean manuscript lines 413–415; detailed results in Supplementary Table S8.'),
('The additional model specifications are described in the Methods (date-effect check, clean manuscript lines 395–399; fixed project effects and covariates, lines 413–415). The country/date alternatives left every habitat contrast’s direction and significance decision unchanged (Supp. Text S3, “Model assumptions”, Table S7). Indonesia exclusion retained directions but previously significant differences became non-significant (clean manuscript lines 446–448; Supplementary Table S3).', 'Alternative-model methods, clean manuscript lines 395–399 and 413–415; Indonesia result, lines 446–448.'),
],
15:[
('We now describe healthy sites as local reference conditions and state that they were not undisturbed controls (clean manuscript line 586).', 'Local-reference limitation, clean manuscript line 586.'),
('We distinguish historical causes of degradation from ongoing exposure. Ongoing stressors at each reference site were not quantified, so we cannot add reliable site-specific descriptions without further field records (clean manuscript lines 587–588). This remains an author follow-up.', 'Ongoing-exposure limitation, clean manuscript lines 587–588.'),
],
16:[
('We now state that the acoustic attribution of Scrape to parrotfish grazing has not been validated by direct observation in every region (clean manuscript lines 371–372).', 'Regional validation limitation, clean manuscript lines 371–372.'),
('We report that 26 of 100 reviewed Mexican detections were false positives and describe this as the proportion among retained detections, not the conventional rate among negative cases (clean manuscript lines 326–331 and 373–374).', 'Precision terminology and Mexico result, clean manuscript lines 326–331 and 373–374.'),
('The available checked detections do not identify the alternative sound sources reliably. We have stated the uncertainty in the Methods and Discussion (clean manuscript lines 373–375 and 637–640); identifying sources requires the author’s review of the recordings.', 'Uncertain sound attribution, clean manuscript lines 373–375 and 637–640.'),
],
19:[
('Thank you. The requested issues are addressed separately below and in responses 10, 13 and 17.', 'No separate manuscript edit for this introductory point.'),
('Excluding Indonesia retained all effect directions, but previously significant contrasts became non-significant (clean manuscript lines 446–449; Supp. Text S3, “Indonesia’s influence”, Table S3). The Discussion qualifies generalisation (lines 527–530).', 'Indonesia sensitivity and qualification, clean manuscript lines 446–449 and 527–530.'),
('We fitted project as a fixed effect and examined available coral cover. The project check changed no direction or significance decision (Supp. Text S3, “Model assumptions”, Table S7). Cover changed some comparisons (Supp. Text S3, “Coral cover”, Table S8). Both checks are described in the Methods (clean manuscript lines 413–415).', 'Project and covariate checks, clean manuscript lines 413–415.'),
('We now identify the single-restored-site comparisons as preliminary and distinguish repeated days from independent sites (clean manuscript lines 641–644). Survey measures are not comparable across all five regions, so the data do not support a measured global richness gradient.', 'Spatial-replication limitation, clean manuscript lines 641–644.'),
('Thank you. The relevant revisions and their locations are given in the responses above.', 'No separate manuscript edit for this closing point.'),
],
22:[
('The draft response and manuscript are prepared for double-anonymous review; the final anonymity check remains on the author TODO list.', 'Submission check; no manuscript edit for this instruction.'),
('Clean and tracked manuscript and supplement files are supplied.', 'Submission files prepared; no manuscript text edit.'),
('Suggested applied Abstract and synthesis wording is in the author TODO section for author approval.', 'No manuscript edit yet.'),
('The graphical abstract remains an author submission task.', 'No manuscript edit yet.'),
('An anonymous data-availability statement remains on the author TODO list.', 'No manuscript edit yet.'),
('The manuscript word limit requires a final check after author edits.', 'No manuscript edit for this instruction.'),
('The supporting analyses are in Supplementary Text S3 and Tables S3–S9.', 'Supplementary material updated; no extra manuscript edit for this instruction.'),
('The return deadline or extension must be checked in the submission system.', 'No manuscript edit for this instruction.'),
('The optional additional-language Abstract remains an author choice.', 'No manuscript edit.'),
('Publication licence clearance remains an author task.', 'No manuscript edit.'),
],
}

TODO_LOCATIONS={
    'A1':'Abstract lines 21–25; Discussion lines 527–530; Conclusion lines 647–652.',
    'A2':'Modelling Methods lines 390–395 and 405–417; sensitivity Results lines 446–450.',
    'A3':'Restoration-site age and category lines 149–151; site identity is in Supplementary Table S1.',
    'A4':'Local-reference and stressor limitation, lines 586–590.',
    'A5':'Mexican Scrape false positives and source uncertainty, lines 373–375.',
    'A6':'Bioturbation claim, lines 83–85; proxy limitation, lines 637–640.',
    'A7':'Figure 2 caption, lines 207–211; Figure 4 caption begins at line 453 and Figure 5 caption at line 509.',
    'A8':'No data-availability statement has been inserted; see the References boundary before submission.',
    'A9':'Restoration methods, lines 131–135; Figure 2 workflow, lines 207–211; bibliography needs separate author review.',
}

def style(doc: Document) -> None:
    doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.sections[0].top_margin=Inches(.7);doc.sections[0].bottom_margin=Inches(.7)
    doc.core_properties.author='';doc.core_properties.last_modified_by=''

def blue(p,text: str):
    r=p.add_run(text);r.font.color.rgb=BLUE;return r

def comment_text(p,text: str,annotate: bool):
    """Keep reviewer text black; add italic blue current-line mappings."""
    pattern=r'\b(?:L|Lines?\s+)(\d+(?:[–-]\d+)?(?:,\s*\d+(?:[–-]\d+)?)*)'
    cursor=0
    for match in re.finditer(pattern,text,flags=re.I):
        p.add_run(text[cursor:match.start()]);p.add_run(match.group(0))
        if annotate:
            ranges=[]
            for bit in match.group(1).split(','):
                nums=[int(x) for x in re.findall(r'\d+',bit)];start=nums[0];end=nums[-1]
                locations=[LINES['original_to_revised'].get(str(n)) for n in range(start,end+1)]
                locations=[v for v in locations if v]
                ranges.append(f'{min(v[0] for v in locations)}–{max(v[1] for v in locations)}' if locations else 'verify in Word')
            r=blue(p,' (now lines '+', '.join(ranges)+')');r.italic=True
        cursor=match.end()
    p.add_run(text[cursor:])

def build(populated: bool) -> None:
    doc=Document();style(doc)
    doc.add_heading('Response to reviewers — working draft' if populated else 'Reviewer comments and revision checklist',0)
    doc.add_paragraph('Stable numbering: Reviewer 1, comments 1–9; Reviewer 2, 10–18; Associate Editor, 19; Editor, 20–21; resubmission instructions, 22. Lettered parts retain the parent number.')
    if populated:
        doc.add_paragraph('Black text reproduces reviewer comments; blue text contains draft responses and precise edit locations. All new manuscript line numbers refer to the accompanying clean-manuscript PDF, Manuscript_line_reference.pdf. Microsoft Word may repaginate. Author decisions and manual tasks are at the bottom; remove that internal section before submission.')
    group=None
    for e in REGISTER['comments']:
        if e['reviewer']!=group:
            group=e['reviewer'];doc.add_heading(group,1)
            if group in REGISTER['preambles']:doc.add_paragraph(REGISTER['preambles'][group])
            if group=='Reviewer 2':doc.add_paragraph(REGISTER['preambles']['Reviewer 2 covering note'])
        doc.add_heading(f'Comment {e["id"]}',2)
        if e.get('parts'):
            for i,part in enumerate(e['parts']):
                p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.18)
                p.add_run(f'{e["id"]}{chr(97+i)}. ').bold=True;comment_text(p,part,False)
                if populated:
                    reply,location=PART_RESPONSES[e['id']][i]
                    blue(doc.add_paragraph(),'Response to '+str(e['id'])+chr(97+i)+': '+reply)
                    blue(doc.add_paragraph(),'Manuscript change: '+location)
        else:
            p=doc.add_paragraph();p.add_run(f'{e["id"]}. ').bold=True;comment_text(p,e['comment'],False)
        if populated and not e.get('parts'):
            status,paragraphs,response=RESPONSES[e['id']]
            blue(doc.add_paragraph(),'Response: '+response)
            blue(doc.add_paragraph(),'Manuscript change: '+MANUSCRIPT_CHANGES[e['id']])
    doc.add_heading('Consolidated revision TODO list',1)
    for ident,refs,task,status in TASKS:
        doc.add_paragraph(f'{ident} — Comments {refs}: {task}'+(f' Status: {status}.' if populated else ''))
    if populated:
        doc.add_heading('AUTHOR TODO — decisions, suggested wording and manual work',1)
        doc.add_paragraph('These suggestions have not been inserted into the manuscript. Current lines refer to the accompanying fixed PDF. Each item restates the issue so you do not need to scroll back through the reviews.')
        for ident,refs,paragraphs,issue,suggestion in TODO:
            doc.add_heading(f'{ident} — Comments {refs}',2)
            doc.add_paragraph(issue)
            doc.add_paragraph('Clean manuscript locations: '+TODO_LOCATIONS[ident])
            for block in suggestion.split('\n'):blue(doc.add_paragraph(),block)
    filename='Responses_populated.docx' if populated else 'Comments_reference.docx'
    doc.save(ROOT/'outputs'/filename)

if __name__=='__main__':
    build(False);build(True)
    (ROOT/'outputs/author_todos.json').write_text(json.dumps(TODO,indent=2,ensure_ascii=False))
    print('Reference and populated response documents written.')
